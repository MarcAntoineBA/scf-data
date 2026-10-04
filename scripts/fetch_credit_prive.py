#!/usr/bin/env python3
"""fetch_credit_prive.py — Données de l'onglet « Crédit Privé ».

Deux couches :
  1. LIVE — séries FRED (api.stlouisfed.org) : délinquances par type de crédit,
     charge-offs, spreads de crédit (HY/BB/BBB/CCC OAS), durcissement bancaire
     (SLOOS), conditions financières (NFCI), prêts des banques aux institutions
     financières non bancaires (canal de contagion), encours de crédit conso,
     ratio de service de la dette des ménages.
  2. STATIQUE SOURCÉE — chiffres du marché du crédit privé institutionnel qui
     n'existent PAS en série temporelle publique (taille de marché, dry powder,
     taux de défaut KBRA/Lincoln, part PIK, non-accruals, décotes BDC), la
     chronologie des effondrements, et le tableau NY Fed des transitions vers
     défaut par type d'emprunt. Chaque donnée porte sa source + URL cliquable.

Sortie : credit_prive_cache.js (window.__CREDIT_PRIVE__ = {...}) + .json
Convention identique aux autres fetchers du site (economie_physique, etc.).
"""
import json
import os
import re
import sys
import time
from datetime import datetime, timezone
from pathlib import Path
from urllib.error import HTTPError, URLError
from urllib.request import Request, urlopen

from _fred_helpers import fetch_fred

HERE = Path(__file__).resolve().parent
# Cibles d'écriture : le cache canonique lu par snapshot_site.sh (~/Library/Caches/
# site_crypto_finance) + le dépôt local (pour servir :8000 immédiatement). Sous
# launchd l'écriture Desktop est bloquée par TCC → best-effort (try/except).
CACHE_DIR = Path.home() / "Library" / "Caches" / "site_crypto_finance"
REPO_DIR = Path(os.path.expanduser("~/Desktop/Site_Crypto_Finance"))
_OUT_DIRS = []
for _d in (CACHE_DIR, REPO_DIR, HERE):
    if _d not in _OUT_DIRS:
        _OUT_DIRS.append(_d)

# ─────────────────────────────────────────────────────────────────────────────
# 1. SÉRIES FRED — (clé, id FRED, libellé FR, start, units)
#    units : 'pct' = %, 'bn' = milliards $, 'idx' = indice/ratio
# ─────────────────────────────────────────────────────────────────────────────
FRED_SERIES = [
    # — Spreads de crédit (quotidien) : le thermomètre du risque —
    ("hy_oas",       "BAMLH0A0HYM2", "Spread High Yield US (OAS)",          "2004-01-01", "pct"),
    ("bb_oas",       "BAMLH0A1HYBB", "Spread BB (OAS)",                     "2004-01-01", "pct"),
    ("bbb_oas",      "BAMLC0A4CBBB", "Spread BBB (OAS)",                    "2004-01-01", "pct"),
    ("ccc_oas",      "BAMLH0A3HYC",  "Spread CCC & moins (OAS)",            "2004-01-01", "pct"),
    # — Délinquances bancaires par type (trimestriel) : « tracker des défauts » —
    ("deliq_consumer", "DRCLACBS",   "Délinquance crédit conso (banques)",  "2000-01-01", "pct"),
    ("deliq_cc",       "DRCCLACBS",  "Délinquance cartes de crédit",        "2000-01-01", "pct"),
    ("deliq_mortgage", "DRSFRMACBS", "Délinquance crédit immobilier",       "2000-01-01", "pct"),
    ("deliq_cre",      "DRCRELEXFACBS","Délinquance immobilier commercial", "2000-01-01", "pct"),
    ("deliq_business", "DRBLACBS",   "Délinquance prêts aux entreprises",   "2000-01-01", "pct"),
    ("chargeoff_cc",   "CORCCACBS",  "Taux de perte cartes de crédit",      "2000-01-01", "pct"),
    # — Ménages : capacité de remboursement —
    ("dsr_household",  "TDSP",       "Service de la dette / revenu (ménages)","2000-01-01","pct"),
    ("consumer_credit","TOTALSL",    "Encours total crédit conso",          "2000-01-01", "bn"),
    ("revolving",      "REVOLSL",    "Crédit renouvelable (cartes)",        "2000-01-01", "bn"),
    # SLOAS (G.19) a été ARRÊTÉE par la Fed au T4 2024. Relève : la même mesure
    # dans les comptes financiers Z.1 (valeurs identiques sur 2006-2024, vérifié
    # le 04/10/2026), toujours publiée chaque trimestre.
    ("student_loans",  "BOGZ1FL153166220Q", "Encours total prêts étudiants (ménages, Z.1)", "2006-01-01", "bn"),
    # — Système : durcissement, conditions, contagion —
    ("sloos_ci",       "DRTSCILM",   "Banques durcissant le crédit aux entreprises", "2000-01-01","pct"),
    ("nfci",           "NFCI",       "Conditions financières (Chicago Fed)","2000-01-01", "idx"),
    ("nbfi_loans",     "LNFACBM027SBOG","Prêts des banques aux institutions financières non bancaires","2015-01-01","bn"),
]


def build_fred():
    out, ok, failed = {}, [], []
    for key, fid, label, start, units in FRED_SERIES:
        r = fetch_fred(fid, start=start)
        if r and r.get("dates"):
            out[key] = {
                "dates": r["dates"], "values": r["values"],
                "label": label, "fred_id": fid, "units": units,
                "source_url": r["source_url"],
                "last": r["values"][-1], "last_date": r["dates"][-1],
                "fetched_at": datetime.now(timezone.utc).isoformat(timespec="seconds"),
            }
            ok.append(f"FRED:{fid}")
            sys.stderr.write(f"[credit_prive] OK {fid:16s} {label[:40]:40s} {r['dates'][-1]} = {r['values'][-1]}\n")
        else:
            failed.append(f"FRED:{fid}")
            sys.stderr.write(f"[credit_prive] FAIL {fid}\n")
        time.sleep(0.25)
    return out, ok, failed


# ─────────────────────────────────────────────────────────────────────────────
# 1 bis. REPLI — on ne perd jamais une donnée déjà collectée.
#    Si une source tombe, on reprend sa DERNIÈRE version réellement collectée
#    (cache précédent), marquée "stale": true avec sa date de collecte d'origine.
#    Jamais de constante réinjectée en silence ; jamais de recul de date.
# ─────────────────────────────────────────────────────────────────────────────
def load_previous():
    best, best_ts = None, ""
    for d in _OUT_DIRS:
        f = d / "credit_prive_cache.json"
        try:
            prev = json.loads(f.read_text(encoding="utf-8"))
        except (OSError, ValueError):
            continue
        ts = (prev.get("meta") or {}).get("updated_at") or ""
        if ts > best_ts:
            best, best_ts = prev, ts
    return best or {}


def _merge_series(new, prev, ident_key):
    """new/prev : dict clé -> série. Garde la plus récente ; reprend l'ancienne si
    la nouvelle manque (même identifiant de source uniquement)."""
    out = dict(new)
    for k, old in (prev or {}).items():
        if not isinstance(old, dict) or not old.get("dates"):
            continue
        cur = out.get(k)
        if cur is None:
            out[k] = dict(old, stale=True)
            sys.stderr.write(f"[credit_prive] REPLI {k} : dernière collecte {old.get('fetched_at', '?')}\n")
        elif cur.get(ident_key) == old.get(ident_key) and str(cur.get("last_date", "")) < str(old.get("last_date", "")):
            out[k] = dict(old, stale=True)   # jamais à reculons
            sys.stderr.write(f"[credit_prive] GARDE {k} : la source a reculé ({cur.get('last_date')} < {old.get('last_date')})\n")
    return out


# ─────────────────────────────────────────────────────────────────────────────
# 1 ter. PRIVATE EQUITY (acte II du chapitre 09) — les 2 seules séries publiques
#    vivantes. Portées par ce collecteur planifié (launchd quotidien + cloud 24 h)
#    pour ne plus dépendre de fetch_pe_data.py (manuel). La page lit d'abord
#    __CREDIT_PRIVE__.pe_live, puis retombe sur __PE_DATA__.series.
# ─────────────────────────────────────────────────────────────────────────────
UA_WEB = "Mozilla/5.0 (Macintosh; Intel Mac OS X 10_15_7) SiteCryptoFinance"


def _get(url, timeout=40, retries=2, binary=False):
    last = None
    for i in range(retries + 1):
        try:
            req = Request(url, headers={"User-Agent": UA_WEB, "Accept": "*/*"})
            with urlopen(req, timeout=timeout) as r:
                body = r.read()
            return body if binary else body.decode("utf-8", "ignore")
        except (HTTPError, URLError, TimeoutError, OSError) as e:
            last = e
            if isinstance(e, HTTPError) and e.code == 404:
                break
            time.sleep(2 * (i + 1))
    sys.stderr.write(f"[credit_prive] échec {url[:90]} : {last}\n")
    return None


def build_pe_live():
    out = {}
    now = datetime.now(timezone.utc).isoformat(timespec="seconds")
    # Banque mondiale CM.MKT.LDOM.NO : sociétés domestiques cotées (US), annuel (source WFE)
    raw = _get("https://api.worldbank.org/v2/country/US/indicator/CM.MKT.LDOM.NO?format=json&per_page=80")
    try:
        d = json.loads(raw) if raw else None
        rows = sorted((int(o["date"]), float(o["value"])) for o in d[1] if o.get("value") is not None)
        if rows:
            out["cotees_us"] = {
                "dates": [f"{y}-12-31" for y, _ in rows], "values": [v for _, v in rows],
                "label": "Sociétés domestiques cotées (US)", "wb_id": "CM.MKT.LDOM.NO",
                "source": "Banque mondiale · CM.MKT.LDOM.NO (source WFE)",
                "source_url": "https://data.worldbank.org/indicator/CM.MKT.LDOM.NO?locations=US",
                "last": rows[-1][1], "last_date": f"{rows[-1][0]}-12-31", "fetched_at": now,
                "wb_lastupdated": (d[0] or {}).get("lastupdated"),
            }
    except (ValueError, KeyError, TypeError, IndexError) as e:
        sys.stderr.write(f"[credit_prive] Banque mondiale illisible : {e}\n")
    # FRED QPRHLDTPEPEQHOLDINGSUSNO : PE des 100 plus grandes pensions publiques (M$ -> Md$)
    r = fetch_fred("QPRHLDTPEPEQHOLDINGSUSNO")
    if r and r.get("values"):
        vals = [v / 1000.0 for v in r["values"]]
        out["pe_pensions"] = {
            "dates": r["dates"], "values": vals, "unit": "Md$",
            "label": "PE détenu par les 100 plus grandes pensions publiques US",
            "fred_id": "QPRHLDTPEPEQHOLDINGSUSNO", "source": "FRED · Census QSPP",
            "source_url": "https://fred.stlouisfed.org/series/QPRHLDTPEPEQHOLDINGSUSNO",
            "last": round(vals[-1], 1), "last_date": r["dates"][-1], "fetched_at": now,
        }
    return out


# ─────────────────────────────────────────────────────────────────────────────
# 1 quater. NY FED — Household Debt and Credit (HHDC), trimestriel, classeur
#    public sans clé. Composition de la dette des ménages, part en retard,
#    impayés 90+ j par type. On prend le DERNIER classeur publié.
# ─────────────────────────────────────────────────────────────────────────────
HHDC_XLSX = "https://www.newyorkfed.org/medialibrary/interactives/householdcredit/data/xls/HHD_C_Report_{y}Q{q}.xlsx"
HHDC_PAGE = "https://www.newyorkfed.org/microeconomics/hhdc"


def _qlabel(s):
    """'26:Q2' -> ('2026-04-01', 'T2 2026')"""
    yy, q = s.split(":Q")
    y = 2000 + int(yy)
    q = int(q)
    return f"{y}-{3 * (q - 1) + 1:02d}-01", f"T{q} {y}"


def build_hhdc():
    try:
        import io
        import openpyxl
    except ImportError as e:
        sys.stderr.write(f"[credit_prive] HHDC : openpyxl absent ({e})\n")
        return None
    now = datetime.now(timezone.utc)
    y, q = now.year, (now.month - 1) // 3 + 1
    for _ in range(6):                      # trimestre courant puis 5 en arrière
        url = HHDC_XLSX.format(y=y, q=q)
        blob = _get(url, binary=True, retries=1)
        if blob and blob[:2] == b"PK":
            break
        q -= 1
        if q == 0:
            y, q = y - 1, 4
    else:
        return None
    try:
        wb = openpyxl.load_workbook(io.BytesIO(blob), read_only=True, data_only=True)

        def rows(sheet):
            return [r for r in wb[sheet].iter_rows(values_only=True)
                    if r and isinstance(r[0], str) and ":Q" in r[0]]

        release = None
        for r in wb["TABLE OF CONTENTS"].iter_rows(values_only=True):
            if r and len(r) > 1 and isinstance(r[1], str) and re.match(r"^[A-Z][a-z]+ \d{4}$", r[1].strip()):
                release = r[1].strip()
                break
        comp = rows("Page 3 Data")          # Mortgage, HE Revolving, Auto, CC, Student, Other, Total (T$)
        dq = rows("Page 11 Data")           # % du solde : Current, 30, 60, 90, 120+, SevDerog
        s90 = rows("Page 12 Data")          # % du solde 90+ par type
        last = comp[-1]
        date, lab = _qlabel(last[0])
        tot = [float(r[7]) for r in comp]
        imax = max(range(len(tot)), key=lambda i: tot[i])
        out = {
            "quarter": lab, "date": date, "release": release,
            "url_xlsx": HHDC_XLSX.format(y=y, q=q), "url": HHDC_PAGE,
            "fetched_at": now.isoformat(timespec="seconds"),
            "total": round(tot[-1], 3),
            "record_total": round(tot[imax], 3), "record_quarter": _qlabel(comp[imax][0])[1],
            "composition": {k: round(float(last[i]), 4) for k, i in
                            (("mortgage", 1), ("heloc", 2), ("auto", 3), ("cards", 4), ("student", 5), ("other", 6))},
            "total_dates": [_qlabel(r[0])[0] for r in comp], "total_values": [round(t, 4) for t in tot],
        }
        if dq and dq[-1][0] == last[0]:
            out["delinquent_pct"] = round(100.0 - float(dq[-1][1]), 2)
        if s90 and s90[-1][0] == last[0]:
            r = s90[-1]
            out["serious90"] = {k: round(float(r[i]), 2) for k, i in
                                (("mortgage", 1), ("heloc", 2), ("auto", 3), ("cards", 4), ("student", 5), ("other", 6), ("all", 7))}
            out["student90_dates"] = [_qlabel(x[0])[0] for x in s90]
            out["student90_values"] = [round(float(x[5]), 2) for x in s90]
        sys.stderr.write(f"[credit_prive] OK NY Fed HHDC {lab} : {out['total']} T$ ({release})\n")
        return out
    except Exception as e:                  # classeur réorganisé : on garde l'ancien
        sys.stderr.write(f"[credit_prive] HHDC illisible ({url}) : {e}\n")
        return None


# ─────────────────────────────────────────────────────────────────────────────
# 2. STATIQUE SOURCÉE — voir bloc STATIC ci-dessous.
#    (rempli depuis le dossier de recherche vérifié — chaque entrée a sa source)
# ─────────────────────────────────────────────────────────────────────────────
from credit_prive_static import STATIC  # noqa: E402


def main():
    prev = load_previous()
    fred, ok, failed = build_fred()
    fred = _merge_series(fred, prev.get("fred"), "fred_id")
    pe_live = _merge_series(build_pe_live(), prev.get("pe_live"), "label")
    for k in ("cotees_us", "pe_pensions"):
        (ok if k in pe_live and not pe_live[k].get("stale") else failed).append(f"PE:{k}")
    hhdc = build_hhdc()
    if hhdc:
        ok.append("NYFED:HHDC")
    else:
        failed.append("NYFED:HHDC")
        old = prev.get("hhdc")
        if old and old.get("total"):
            hhdc = dict(old, stale=True)
            sys.stderr.write(f"[credit_prive] REPLI HHDC : {old.get('quarter')} (collecté {old.get('fetched_at')})\n")
    now = datetime.now(timezone.utc)
    payload = {
        "meta": {
            "updated_at": now.isoformat(),
            "updated_at_unix": int(now.timestamp()),
            "sources_ok": ok,
            "sources_failed": failed,
            "doc_version": "1.0",
        },
        "fred": fred,
        "pe_live": pe_live,
    }
    if hhdc:
        payload["hhdc"] = hhdc
    payload.update(STATIC)

    blob_json = json.dumps(payload, ensure_ascii=False)
    js = ("/* credit_prive_cache.js — généré " + now.isoformat() + " */\n"
          "window.__CREDIT_PRIVE__ = " + blob_json + ";\n")
    written = []
    for d in _OUT_DIRS:
        try:
            d.mkdir(parents=True, exist_ok=True)
            (d / "credit_prive_cache.js").write_text(js, encoding="utf-8")
            (d / "credit_prive_cache.json").write_text(blob_json, encoding="utf-8")
            written.append(str(d))
        except (OSError, PermissionError) as e:
            sys.stderr.write(f"[credit_prive] skip {d}: {e}\n")
    sys.stderr.write(f"[credit_prive] écrit credit_prive_cache.js/.json dans {len(written)} cible(s) "
                     f"({len(ok)} séries OK, {len(failed)} échecs)\n")


if __name__ == "__main__":
    main()
